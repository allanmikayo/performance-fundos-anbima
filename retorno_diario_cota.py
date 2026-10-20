"""Retorno DIÁRIO da cota, por grupo (padrão: infra), ponderado por PL.

Entrega ENXUTA: um Excel pequeno só com a série  data | retorno_pct  (mais
nº de fundos e um índice base 100 de cortesia). Todo o cálculo é feito em
Python; o arquivo NÃO carrega as colunas auxiliares nem fórmulas, para ficar
leve. Se quiser auditar, ligue EXPORTAR_TRILHA e sai um CSV com tudo.

Usa SÓ as cotas e o PL do serie_historica.csv (captação você cruza de outra
base depois).

Filtros/tratamentos:
  1) GRUPO: só os fundos do grupo alvo (classificacao_fundos.csv).
  2) REPORT DIÁRIO: só fundos com mediana do intervalo entre cotas
     <= MAX_GAP_MEDIANO_DIAS e no mínimo MIN_OBS cotas.
  3) BURACOS/comprimentos desiguais: retorno entre observações consecutivas da
     cota, válido só se o gap até a anterior <= MAX_GAP_RETORNO_DIAS; o peso é
     renormalizado a cada dia entre os fundos presentes.

Fórmulas (calculadas em Python):
  retorno  = cota / cota_ant - 1
  peso     = pl_ant / soma(pl_ant dos válidos no dia)
  retorno_pct (do grupo) = soma(peso * retorno)
  indice_t = indice_(t-1) * (1 + retorno_pct_t),  base 100

SAÍDA: analise_output/retorno_diario_cota.xlsx  (aba performance_diaria + metodologia)
       analise_output/retorno_diario_trilha.csv (só se EXPORTAR_TRILHA = True)
"""
from __future__ import annotations

# ╔══════════════════════════════════════════════════════════════════╗
# ║                         CONFIGURAÇÃO                              ║
# ╚══════════════════════════════════════════════════════════════════╝
from pathlib import Path

PASTA_SCRIPT = Path(__file__).resolve().parent

ARQUIVO_SERIE = None                                   # None = acha o mais recente em output/
ARQUIVO_CLASSIFICACAO = PASTA_SCRIPT / "classificacao_fundos.csv"
COL_CHAVE_CLASSIF = "chave"
COL_GRUPO = "grupo"
CHAVE_SERIE = "codigo_classe_consultado"               # ou "cnpj_input"
GRUPO_ALVO = "infra"

MAX_GAP_MEDIANO_DIAS = 4     # define "fundo diário"
MIN_OBS = 30                 # ignora séries curtas
MAX_GAP_RETORNO_DIAS = 5     # retorno só vale se a cota anterior está a <= N dias

EXPORTAR_TRILHA = False      # True => grava também um CSV com as colunas auxiliares (cota_ant, peso, etc.)
INDICE_BASE = 100.0
PASTA_SAIDA = PASTA_SCRIPT / "analise_output"
# ════════════════════ FIM DA CONFIGURAÇÃO ═══════════════════════════


import sys
from datetime import datetime

import pandas as pd

LOG: list[str] = []


def log(msg: str) -> None:
    linha = f"[{datetime.now():%H:%M:%S}] {msg}"
    print(linha, flush=True)
    LOG.append(linha)


def achar_serie() -> Path:
    if ARQUIVO_SERIE:
        p = Path(ARQUIVO_SERIE)
        if not p.exists():
            sys.exit(f"ARQUIVO_SERIE não encontrado: {p}")
        return p
    cand = sorted(PASTA_SCRIPT.glob("output/*/serie_historica.csv"))
    if not cand:
        sys.exit("Não achei output/*/serie_historica.csv. Rode a extração ou preencha ARQUIVO_SERIE.")
    return cand[-1]


def ler_serie(p: Path) -> pd.DataFrame:
    log(f"Lendo série: {p}")
    df = pd.read_csv(p, sep=";", dtype=str, encoding="utf-8-sig", keep_default_na=False)
    for c in [CHAVE_SERIE, "data_competencia", "valor_cota", "valor_patrimonio_liquido"]:
        if c not in df.columns:
            sys.exit(f"Coluna ausente na série: {c}. Achei: {list(df.columns)}")

    def num(s):
        return pd.to_numeric(s.str.replace(".", "", regex=False).str.replace(",", ".", regex=False),
                             errors="coerce")

    out = pd.DataFrame({
        "chave": df[CHAVE_SERIE].str.strip(),
        "data": pd.to_datetime(df["data_competencia"], errors="coerce"),
        "cota": num(df["valor_cota"]),
        "pl": num(df["valor_patrimonio_liquido"]),
    }).dropna()
    out = out[(out["cota"] > 0) & (out["pl"] > 0) & (out["chave"] != "")]
    out = out.drop_duplicates(["chave", "data"]).sort_values(["chave", "data"])
    log(f"  linhas válidas: {len(out)}; fundos: {out['chave'].nunique()}; "
        f"datas: {out['data'].min():%Y-%m-%d} a {out['data'].max():%Y-%m-%d}")
    return out


def ler_classificacao() -> pd.DataFrame:
    if not Path(ARQUIVO_CLASSIFICACAO).exists():
        sys.exit(f"Classificação não encontrada: {ARQUIVO_CLASSIFICACAO}. "
                 "Rode a outra análise uma vez para gerar a fictícia, ou crie a real.")
    cab = Path(ARQUIVO_CLASSIFICACAO).read_text(encoding="utf-8-sig").splitlines()[0]
    sep = ";" if ";" in cab else ","
    cls = pd.read_csv(ARQUIVO_CLASSIFICACAO, dtype=str, sep=sep, encoding="utf-8-sig")
    if COL_CHAVE_CLASSIF not in cls or COL_GRUPO not in cls:
        sys.exit(f"Classificação precisa das colunas '{COL_CHAVE_CLASSIF}' e '{COL_GRUPO}'.")
    cls = cls.rename(columns={COL_CHAVE_CLASSIF: "chave", COL_GRUPO: "grupo"})
    cls["chave"] = cls["chave"].astype(str).str.strip()
    cls["grupo"] = cls["grupo"].astype(str).str.strip()
    return cls[cls["chave"] != ""].drop_duplicates("chave")


def selecionar_diarios(serie: pd.DataFrame) -> tuple[set, pd.DataFrame]:
    """Fundos diários = mediana do intervalo entre cotas <= MAX_GAP_MEDIANO_DIAS e n_obs >= MIN_OBS."""
    def stats(g):
        d = g["data"].sort_values()
        return pd.Series({"n_obs": len(d), "gap_mediano_dias": d.diff().dt.days.median()})
    freq = serie.groupby("chave").apply(stats, include_groups=False).reset_index()
    ok = freq[(freq["gap_mediano_dias"] <= MAX_GAP_MEDIANO_DIAS) & (freq["n_obs"] >= MIN_OBS)]
    return set(ok["chave"]), freq


def calcular(obs: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Colunas auxiliares (em Python) + série diária do grupo."""
    o = obs.sort_values(["chave", "data"]).copy()
    o["cota_ant"] = o.groupby("chave")["cota"].shift(1)
    o["pl_ant"] = o.groupby("chave")["pl"].shift(1)
    o["gap_dias"] = o.groupby("chave")["data"].diff().dt.days
    o["retorno"] = o["cota"] / o["cota_ant"] - 1
    o["valido"] = (o["retorno"].notna() & (o["gap_dias"] <= MAX_GAP_RETORNO_DIAS)
                   & (o["pl_ant"] > 0) & (o["cota_ant"] > 0))

    v = o[o["valido"]].copy()
    v["soma_pl_dia"] = v.groupby("data")["pl_ant"].transform("sum")   # Σ PL_ant dos válidos no dia
    v["peso"] = v["pl_ant"] / v["soma_pl_dia"]                        # peso ponderado por PL
    v["contribuicao"] = v["peso"] * v["retorno"]                     # peso * retorno

    perf = (v.groupby("data")
              .agg(retorno_pct=("contribuicao", "sum"), n_fundos=("chave", "nunique"))
              .reset_index())
    perf["indice_base100"] = INDICE_BASE * (1 + perf["retorno_pct"]).cumprod()

    # devolve também as colunas auxiliares (para a trilha opcional)
    o = o.merge(v[["data", "chave", "soma_pl_dia", "peso", "contribuicao"]],
                on=["data", "chave"], how="left")
    return o, perf


def exportar(perf: pd.DataFrame, saida: Path) -> None:
    from openpyxl.styles import numbers
    with pd.ExcelWriter(saida, engine="openpyxl") as w:
        perf.to_excel(w, sheet_name="performance_diaria", index=False)
        ws = w.sheets["performance_diaria"]
        for linha in range(2, len(perf) + 2):
            ws.cell(linha, 1).number_format = "yyyy-mm-dd"      # data
            ws.cell(linha, 2).number_format = "0.0000%"         # retorno_pct
            ws.cell(linha, 4).number_format = "0.0000"          # indice
        met = pd.DataFrame({"metodologia": [
            f"Grupo: {GRUPO_ALVO}. So fundos diarios (gap mediano <= {MAX_GAP_MEDIANO_DIAS}, obs >= {MIN_OBS}).",
            "retorno do fundo = cota / cota_ant - 1 (cota_ant = cota da observacao anterior do fundo).",
            f"so vale como diario se o gap ate a cota anterior <= {MAX_GAP_RETORNO_DIAS} dias.",
            "peso = pl_ant / soma(pl_ant dos validos no dia); renormalizado por dia.",
            "retorno_pct (grupo) = soma(peso * retorno). indice_t = indice_(t-1) * (1 + retorno_pct_t).",
            "ATENCAO: assume cota = total return; fundos de infra que distribuem subestimam o retorno.",
            "Colunas auxiliares (cota_ant, peso, contribuicao, etc.) saem em CSV se EXPORTAR_TRILHA=True.",
        ]})
        met.to_excel(w, sheet_name="metodologia", index=False)


def main() -> None:
    PASTA_SAIDA.mkdir(parents=True, exist_ok=True)
    serie = ler_serie(achar_serie()).merge(ler_classificacao(), on="chave", how="inner")
    serie = serie[serie["grupo"] == GRUPO_ALVO]
    if serie.empty:
        sys.exit(f"Nenhum fundo do grupo '{GRUPO_ALVO}' casou com a série. Confira CHAVE_SERIE/classificação.")
    log(f"Fundos do grupo '{GRUPO_ALVO}' com série: {serie['chave'].nunique()}")

    chaves, freq = selecionar_diarios(serie)
    log(f"  diários (gap mediano <= {MAX_GAP_MEDIANO_DIAS}, obs >= {MIN_OBS}): {len(chaves)} de {len(freq)}")
    obs = serie[serie["chave"].isin(chaves)][["data", "chave", "cota", "pl"]]
    if obs.empty:
        sys.exit("Nenhum fundo diário após os filtros -- afrouxe MAX_GAP_MEDIANO_DIAS/MIN_OBS.")

    trilha, perf = calcular(obs)

    saida = PASTA_SAIDA / "retorno_diario_cota.xlsx"
    exportar(perf, saida)
    log(f"Excel enxuto salvo: {saida}  ({len(perf)} dias)")
    log(f"  retorno acumulado do grupo = {perf['indice_base100'].iloc[-1]/INDICE_BASE-1:.4%}; "
        f"fundos/dia médio = {perf['n_fundos'].mean():.1f}")

    if EXPORTAR_TRILHA:
        csv = PASTA_SAIDA / "retorno_diario_trilha.csv"
        trilha.to_csv(csv, index=False, sep=";", decimal=",", encoding="utf-8-sig")
        log(f"  trilha completa (colunas auxiliares): {csv}")

    (PASTA_SAIDA / "log_retorno_diario.txt").write_text("\n".join(LOG), encoding="utf-8")


if __name__ == "__main__":
    main()
